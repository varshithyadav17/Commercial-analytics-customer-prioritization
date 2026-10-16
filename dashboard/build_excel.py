"""Build a multi-sheet Excel report from generated model outputs."""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
MODEL = ROOT / "data" / "model_outputs"
OUT = ROOT / "Commercial_Analytics_Report.xlsx"


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in ws.columns:
        values = [str(cell.value or "") for cell in col[:200]]
        width = min(max(max((len(value) for value in values), default=10) + 2, 12), 35)
        ws.column_dimensions[get_column_letter(col[0].column)].width = width


def add_frame(wb, title, frame):
    ws = wb.create_sheet(title)
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        ws.append([None if pd.isna(value) else value.item() if hasattr(value, "item") else value for value in row])
    style_sheet(ws)
    if ws.max_row > 1 and ws.max_column > 0:
        table = Table(displayName=f"T_{len(wb.sheetnames)}", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                              showColumnStripes=False, showFirstColumn=False,
                                              showLastColumn=False)
        ws.add_table(table)
    return ws


def main():
    required = {
        "segment_summary.csv": "Segment Summary",
        "optimized_call_plan.csv": "Contact Plan",
        "icm_payouts.csv": "Quota & Incentives",
        "territory_summary.csv": "Territory Summary",
        "monthly_trend.csv": "Monthly Trend",
    }
    missing = [name for name in required if not (MODEL / name).is_file()]
    if missing:
        raise FileNotFoundError("Run python -m src.run_pipeline first; missing: " + ", ".join(missing))
    wb = Workbook()
    summary = wb.active
    summary.title = "Executive Summary"
    summary.append(["Metric", "Value", "Interpretation"])
    segment = pd.read_csv(MODEL / "segment_summary.csv")
    plan = pd.read_csv(MODEL / "optimized_call_plan.csv")
    icm = pd.read_csv(MODEL / "icm_payouts.csv")
    summary_rows = [
        ("HCP count", int(segment["hcp_count"].sum()), "Unique synthetic HCPs in this run"),
        ("Simulated annual volume", round(float(segment["total_rx_volume"].sum()), 1), "Synthetic proxy; not real prescription data"),
        ("Observed calls", int(segment["total_calls"].sum()), "Total rows in the call activity log"),
        ("Planned monthly calls", round(float(plan["optimized_calls_per_month"].sum()), 2), "Heuristic allocation under stated capacity"),
        ("Illustrative incentive payout", round(float(icm["incentive_payout_usd"].sum()), 0), "Teaching assumption, not payroll"),
    ]
    for row in summary_rows:
        summary.append(list(row))
    style_sheet(summary)
    for filename, title in required.items():
        add_frame(wb, title, pd.read_csv(MODEL / filename))
    methodology = wb.create_sheet("Methodology")
    for row in [
        ["Methodology note", "Interpretation"],
        ["Segmentation", "Rule-based potential and receptivity thresholds; review thresholds before use."],
        ["Contact plan", "Minimum coverage floor plus priority-weighted remaining capacity; a heuristic."],
        ["Quota preview", "Illustrative tiers: <80% = 0%, 80–99% = 50%, 100–109% = 100%, >=110% = 150%."],
        ["Causality", "Synthetic data encodes a contact-volume relationship. It does not identify causal lift."],
        ["Data", "All records are seeded synthetic examples and represent no real healthcare professional or company."],
    ]:
        methodology.append(row)
    style_sheet(methodology)
    wb.save(OUT)
    print(f"Workbook created: {OUT}")


if __name__ == "__main__":
    main()
